"""M22-TRADE-XLSX — trade raporlarını XLSX'e aktar (SALT RENDERER, 2026-09-18).

Kaynak: `reports/trades_*.csv` (capped P0xA0 backtest kayıtları — YENİ analiz YOK).
Çıktı: `reports/xlsx/trades_<asset>.xlsx` · sayfalar: TRADES · OZET · EGRI.

Biçim (direktif M22):
  * başlık satırı zemin renkli + freeze panes + auto-filter + sütun genişlikleri
  * r_multiple koşullu renk: pozitif yeşil / negatif kırmızı (CellIsRule)
  * exit_reason renk kodu: take_profit yeşil · stop kırmızı · time_stop amber
    · end_of_data gri
  * sayısal sütunlar 2-4 ondalık; zamanlar `yyyy-mm-dd hh:mm`
  * OZET: trade sayısı · exit_reason dağılımı · ort/medyan R · win rate ·
    maks ardışık kayıp · cost_total toplamı · net expectancy · equity maxDD
  * EGRI: trade sırasıyla equity + openpyxl çizgi grafik

Sınırlar: yeni analiz YOK (salt renderer); hüküm/seçim/look etkisi YOK;
watch_only — bu dosya tavsiye ÜRETMEZ, yalnız kayıtları biçimlendirir.

Çalıştırma:
    python3 -W ignore scripts/export_trade_report_xlsx.py
"""
from __future__ import annotations

import os
import sys
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(_HERE, "..")
OUT_DIR = os.path.join(ROOT, "reports", "xlsx")
ASSETS = ("BTC", "GOLD", "SILVER", "bist30_basket")
START_EQ = 100_000.0          # Aşama 6 başlangıç sermayesi (sabit kayıt)

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(color="FFFFFF", bold=True)
FILL_TP = PatternFill("solid", fgColor="C6EFCE")
FILL_STOP = PatternFill("solid", fgColor="FFC7CE")
FILL_TS = PatternFill("solid", fgColor="FFEB9C")
FILL_EOD = PatternFill("solid", fgColor="D9D9D9")
REASON_FILL = {"take_profit": FILL_TP, "stop": FILL_STOP,
               "time_stop": FILL_TS, "end_of_data": FILL_EOD}
FMT2 = "0.00"
FMT3 = "0.000"
FMT4 = "0.0000"
TIME_FMT = "yyyy-mm-dd hh:mm"
COL_FMT = {"entry_price": FMT2, "stop_price": FMT2, "tp_price": FMT2, "exit_price": FMT2,
           "units": FMT2, "notional_entry": FMT2, "notional_exit": FMT2,
           "cost_entry": FMT2, "cost_exit": FMT2, "cost_total": FMT2,
           "gross_pnl": FMT2, "net_pnl": FMT2, "r_multiple": FMT3,
           "return_frac": FMT4, "atr_at_entry": FMT2, "stop_distance_atr": FMT3,
           "equity_before": FMT2, "equity_after": FMT2,
           "partial_exit_price": FMT2, "partial_units": FMT2,
           "partial_gross": FMT2, "partial_net": FMT2}
TIME_COLS = ("signal_timestamp_utc", "entry_timestamp_utc", "exit_timestamp_utc",
             "execution_timestamp_utc")
WIDTHS = {"direction": 10, "signal_bar": 11, "signal_timestamp_utc": 18, "entry_bar": 10,
          "entry_timestamp_utc": 18, "entry_price": 12, "stop_price": 12, "tp_price": 12,
          "exit_bar": 9, "exit_timestamp_utc": 18, "exit_price": 12, "exit_reason": 13,
          "bars_held": 10, "units": 10, "notional_entry": 14, "notional_exit": 14,
          "cost_entry": 11, "cost_exit": 11, "cost_total": 11, "gross_pnl": 12,
          "net_pnl": 12, "r_multiple": 11, "equity_before": 13, "equity_after": 13,
          "return_frac": 11, "atr_at_entry": 11, "stop_distance_atr": 15,
          "profile": 9, "group_key": 11}


def _ts(v):
    """Excel tz-aware datetime kabul etmez -> naive UTC'ye normalize et."""
    try:
        t = pd.Timestamp(v)
        if t.tzinfo is not None:
            t = t.tz_convert("UTC").tz_localize(None)
        return t.to_pydatetime()
    except Exception:                                                # noqa: BLE001
        return None


def _max_consec_loss(df: pd.DataFrame) -> int:
    cur = best = 0
    for v in df["net_pnl"]:
        cur = cur + 1 if v < 0 else 0
        best = max(best, cur)
    return best


def _equity_curve(df: pd.DataFrame) -> pd.Series:
    eq = START_EQ + df["net_pnl"].cumsum()
    return pd.concat([pd.Series([START_EQ]), eq], ignore_index=True)


def _summary_rows(df: pd.DataFrame) -> list:
    eq = _equity_curve(df)
    dd = float((eq / eq.cummax() - 1.0).min())
    rows = [
        ("trade sayısı", int(len(df))),
        ("win rate", round(float((df["net_pnl"] > 0).mean()), 4)),
        ("ortalama R", round(float(df["r_multiple"].mean()), 4)),
        ("medyan R", round(float(df["r_multiple"].median()), 4)),
        ("net expectancy (R)", round(float(df["r_multiple"].mean()), 4)),
        ("maks ardışık kayıp", int(_max_consec_loss(df))),
        ("cost_total toplamı", round(float(df["cost_total"].sum()), 2)),
        ("equity maxDD", round(dd, 6)),
        ("başlangıç equity", START_EQ),
        ("bitiş equity", round(float(eq.iloc[-1]), 2)),
    ]
    return rows


def build_workbook(df: pd.DataFrame) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "TRADES"
    cols = list(df.columns)
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(c, 12)
    ri_col = cols.index("r_multiple") + 1 if "r_multiple" in cols else None
    er_col = cols.index("exit_reason") + 1 if "exit_reason" in cols else None
    for i, (_, row) in enumerate(df.iterrows(), start=2):
        for j, c in enumerate(cols, start=1):
            v = row[c]
            if c in TIME_COLS:
                v = _ts(v)
                cell = ws.cell(row=i, column=j, value=v)
                cell.number_format = TIME_FMT
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                cell = ws.cell(row=i, column=j, value=float(v))
                cell.number_format = COL_FMT.get(c, FMT2)
            else:
                cell = ws.cell(row=i, column=j, value=v)
            if c == "exit_reason" and v in REASON_FILL:
                cell.fill = REASON_FILL[v]
    n = len(df) + 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n}"
    if ri_col and n > 2:
        rng = f"{get_column_letter(ri_col)}2:{get_column_letter(ri_col)}{n}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan",
                                      formula=["0"], fill=FILL_TP))
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan",
                                      formula=["0"], fill=FILL_STOP))

    ws2 = wb.create_sheet("OZET")
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 16
    ws2.column_dimensions["C"].width = 10
    h = ws2.cell(row=1, column=1, value="özet")
    h.fill = HDR_FILL
    h.font = HDR_FONT
    r = 2
    for k, v in _summary_rows(df):
        ws2.cell(row=r, column=1, value=k)
        c = ws2.cell(row=r, column=2, value=v)
        c.number_format = FMT4 if isinstance(v, float) else "0"
        r += 1
    r += 1
    h2 = ws2.cell(row=r, column=1, value="exit_reason dağılımı")
    h2.fill = HDR_FILL
    h2.font = HDR_FONT
    r += 1
    for reason, cnt in df["exit_reason"].value_counts().items():
        c1 = ws2.cell(row=r, column=1, value=str(reason))
        c2 = ws2.cell(row=r, column=2, value=int(cnt))
        if reason in REASON_FILL:
            c1.fill = REASON_FILL[reason]
        r += 1

    ws3 = wb.create_sheet("EGRI")
    ws3.column_dimensions["A"].width = 20
    ws3.column_dimensions["B"].width = 14
    ws3.cell(row=1, column=1, value="exit_timestamp_utc").fill = HDR_FILL
    ws3.cell(row=1, column=1).font = HDR_FONT
    hb = ws3.cell(row=1, column=2, value="equity")
    hb.fill = HDR_FILL
    hb.font = HDR_FONT
    eq = _equity_curve(df)
    ws3.cell(row=2, column=1, value=df["entry_timestamp_utc"].iloc[0] if len(df) else None)
    ws3.cell(row=2, column=1).number_format = TIME_FMT
    ws3.cell(row=2, column=2, value=float(eq.iloc[0])).number_format = FMT2
    for i, (_, row) in enumerate(df.iterrows(), start=3):
        c1 = ws3.cell(row=i, column=1, value=_ts(row["exit_timestamp_utc"]))
        c1.number_format = TIME_FMT
        ws3.cell(row=i, column=2, value=float(eq.iloc[i - 2])).number_format = FMT2
    ch = LineChart()
    ch.title = "equity (trade sırası, maliyet dahil)"
    ch.height, ch.width = 8, 22
    data = Reference(ws3, min_col=2, min_row=1, max_row=len(df) + 2)
    cats = Reference(ws3, min_col=1, min_row=2, max_row=len(df) + 2)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.legend = None
    ws3.add_chart(ch, "D2")
    return wb


def main() -> int:
    os.makedirs(OUT_DIR, exist_ok=True)
    ok = 0
    for asset in ASSETS:
        src = os.path.join(ROOT, "reports", f"trades_{asset}.csv")
        if not os.path.exists(src):
            print(f"  [{asset}] CSV yok — atlandı")
            continue
        df = pd.read_csv(src)
        if not len(df):
            print(f"  [{asset}] boş CSV — atlandı")
            continue
        wb = build_workbook(df)
        out = os.path.join(OUT_DIR, f"trades_{asset}.xlsx")
        wb.save(out)
        ok += 1
        print(f"  [{asset}] {len(df)} trade -> {os.path.relpath(out, ROOT)}")
    print(f"\nxlsx raporları: {ok} dosya (reports/xlsx/)")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
